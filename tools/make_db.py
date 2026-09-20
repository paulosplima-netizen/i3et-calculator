"""Build data/base_referencia.sqlite and data/csv/ from the source workbook.

The workbook (Documento 3) is the human-editable form of the base; the SQLite
file is the form the program reads. This script is the only bridge between
them, so the conversion is auditable and repeatable.

Usage:
    python tools/make_db.py [caminho/para/BaseDeDados_CalculadoraBP_AAAAMMDDx.xlsx]
"""

from __future__ import annotations

import csv
import datetime as dt
import glob
import hashlib
import os
import sqlite3
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from core import schema  # noqa: E402

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(REPO, "data", "base_referencia.sqlite")
CSV_DIR = os.path.join(REPO, "data", "csv")

#: Onde procurar a base, quando o caminho nao e passado na linha de comando.
#: Dentro de cada pasta vale a versao mais recente pelo nome, que e o que a
#: convencao _AAAAMMDDx garante ser a ordem cronologica (diretriz D6).
DEFAULT_DIRS = [
    os.path.expanduser(
        "~/mnt/_Calculadora da Pegada de Carbono do Berço ao Portão/Documentação"),
    os.path.expanduser(
        "~/OneDrive/Documents/_Unicamp/_Calculadora da Pegada de Carbono do Berço ao "
        "Portão/Documentação"),
]
SOURCE_GLOB = "BaseDeDados_CalculadoraBP_*.xlsx"

TEXT_ID_COLUMNS = {"IDM"}          # codes such as 4a, 4b, 46a, 46b


def find_source(argv) -> str:
    if len(argv) > 1:
        return argv[1]
    candidates = []
    for d in DEFAULT_DIRS:
        candidates += sorted(glob.glob(os.path.join(d, SOURCE_GLOB)), reverse=True)
    for p in candidates:
        if os.path.exists(p):
            return p
    raise SystemExit("Workbook not found. Pass the path as an argument.")


def cell(value, column: str):
    """Normalise one spreadsheet cell for storage."""
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    if column in TEXT_ID_COLUMNS:
        if isinstance(value, float) and value == int(value):
            value = int(value)
        return str(value).strip()
    if isinstance(value, str):
        return value.strip()
    return value


def read_sheet(wb, name):
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    header = [h for h in next(rows) if h is not None]
    out = []
    for r in rows:
        if r is None or all(v is None for v in r[:len(header)]):
            continue
        out.append({h: cell(v, h) for h, v in zip(header, r)})
    return header, out


def main() -> int:
    src = find_source(sys.argv)
    print("fonte:", src)
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)

    os.makedirs(os.path.dirname(DB), exist_ok=True)
    os.makedirs(CSV_DIR, exist_ok=True)
    if os.path.exists(DB):
        os.remove(DB)

    conn = sqlite3.connect(DB)
    schema.create_schema(conn)
    conn.execute("PRAGMA foreign_keys = ON")

    counts, problems = {}, []
    for sheet, table in schema.SHEET_TO_TABLE.items():
        if sheet not in wb.sheetnames:
            problems.append(f"{sheet}: aba ausente no workbook")
            continue
        header, rows = read_sheet(wb, sheet)
        cols = ", ".join(header)
        marks = ", ".join("?" * len(header))
        sql = f"INSERT INTO {table} ({cols}) VALUES ({marks})"
        for r in rows:
            try:
                conn.execute(sql, [r.get(h) for h in header])
            except sqlite3.Error as exc:
                problems.append(f"{table}: {exc} | linha {r}")
        counts[table] = len(rows)
        with open(os.path.join(CSV_DIR, f"{sheet}.csv"), "w", newline="",
                  encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=header, delimiter=";")
            w.writeheader()
            w.writerows(rows)
    conn.commit()

    fk = conn.execute("PRAGMA foreign_key_check").fetchall()
    integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]

    # Fingerprint of the reference base: identifies which parameters produced a result.
    h = hashlib.sha256()
    for table in schema.PARAM_TABLES:
        for row in conn.execute(f"SELECT * FROM {table}"):
            h.update(repr(row).encode("utf-8"))
    base_hash = h.hexdigest()
    conn.execute("CREATE TABLE IF NOT EXISTS M00_BaseInfo "
                 "(Key TEXT PRIMARY KEY, Value TEXT)")
    conn.executemany("INSERT OR REPLACE INTO M00_BaseInfo VALUES (?,?)", [
        ("ParamBaseHash", base_hash),
        ("SourceWorkbook", os.path.basename(src)),
        ("BuiltAt", dt.datetime.now().isoformat(timespec="seconds")),
        ("SchemaVersion", "1"),
    ])
    conn.commit()
    conn.close()

    total = sum(counts.values())
    print(f"\n{len(counts)} tabelas, {total} linhas")
    for t in schema.PARAM_TABLES + schema.SCENARIO_TABLES:
        print(f"   {t:32s} {counts.get(t, 0):6d}")
    print(f"\nintegrity_check: {integrity}")
    print(f"foreign_key_check: {len(fk)} violacoes")
    for v in fk[:15]:
        print("   ", v)
    if problems:
        by_table = {}
        for p in problems:
            by_table.setdefault(p.split(":")[0], []).append(p)
        print(f"\n{len(problems)} problemas de insercao, por tabela:")
        for tab, items in sorted(by_table.items()):
            print(f"   {tab}: {len(items)}")
            print("      ex.:", items[0][:160])
    print(f"\nParamBaseHash: {base_hash[:16]}...")
    print("gravado:", DB)
    print("csv em :", CSV_DIR)
    return 1 if (fk or problems or integrity != "ok") else 0


if __name__ == "__main__":
    raise SystemExit(main())
